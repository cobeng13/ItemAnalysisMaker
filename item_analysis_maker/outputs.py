from pathlib import Path
from copy import copy
import re
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font
from .analysis import Analysis
from .programs import PROGRAM_LONG_NAMES

class OutputError(RuntimeError): pass

SUMMARY_TEMPLATE = Path(__file__).resolve().parents[1] / "ITEM_ANALYSIS_SUMMARY_FORMAT.docx"

def safe_part(value,fallback):
    return (re.sub(r"[^A-Za-z0-9_-]+","_",value.strip()).strip("_-")[:60] or fallback)

def _shift_merges(ws,at,amount):
    old=list(ws.merged_cells.ranges)
    for r in old: ws.unmerge_cells(str(r))
    for r in old:
        c1,r1,c2,r2=r.bounds
        if r1>=at: r1+=amount; r2+=amount
        elif r2>=at: raise OutputError("A merged template cell crosses the student rows.")
        if c1!=c2 or r1!=r2: ws.merge_cells(start_row=r1,start_column=c1,end_row=r2,end_column=c2)

def _copy_style(ws,source,target):
    for col in range(1,103):
        a,b=ws.cell(source,col),ws.cell(target,col)
        if a.has_style: b._style=a._style
        b.alignment=copy(a.alignment)
    ws.row_dimensions[target].height=ws.row_dimensions[source].height

def _expand(ws,et,el):
    if et:
        _shift_merges(ws,34,et); ws.insert_rows(34,et)
        for r in range(34,34+et): _copy_style(ws,33,r)
    base=61+et
    if el:
        _shift_merges(ws,base,el); ws.insert_rows(base,el)
        for r in range(base,base+el): _copy_style(ws,base-1,r)
    ws.print_area=f"A1:CX{78+et+el}"

def _metadata(ws,m,n):
    labels={"Prelims":"Prelim","Midterms":"Midterm","Finals":"Finals"}
    for row,label in ((10,"Prelim"),(11,"Midterm"),(12,"Finals")):
        ws.cell(row,1).value=("X" if labels[m["exam_type"]]==label else " ")+" "+label
    ws["E10"]=m["academic_year"]; ws["E11"]=m["semester"]
    ws["J10"]=m["subject"]; ws["J11"]=m["description"]; ws["X10"]=n
    ws["X11"]="=X10*0.27"; ws["X14"]="=MIN(ROUNDUP(X11,0),16)"
    ws["A72"]=f"Prepared by: {m['prepared_by']}"
    ws["V72"]=f"Approved by: {m['department_chairperson']}"

def _student_header(ws,row):
    left=f"A{row}:B{row+1}"
    label=ws.cell(row,1)
    style=label._style
    alignment=copy(label.alignment)
    ws.unmerge_cells(left)
    label=ws.cell(row,1)
    target=ws.cell(row,2)
    label.value=None
    target._style=style
    target.alignment=alignment
    target.value="Student Names"
    ws.merge_cells(start_row=row,start_column=2,end_row=row+1,end_column=2)

def _group(ws,records,start,slots,nitems,first_number):
    for i in range(slots):
        row=start+i
        number=first_number+i
        ws.cell(row,1).value=number
        ws.cell(row,2).value=f"Student {number}" if i<len(records) else None
        for col in (1,2):
            cell=ws.cell(row,col)
            old=cell.font
            cell.font=Font(name="Arial Narrow",size=10,bold=old.bold,italic=old.italic,
                           vertAlign=old.vertAlign,underline=old.underline,strike=old.strike,color=old.color)
        for q in range(nitems):
            ws.cell(row,q+3).value=int(records[i].answers[q]) if i<len(records) else None

def _formulas(ws,nitems,hs,hn,ls,ln):
    ht=hs+hn; lt=ls+ln; ph=ht+1; pl=lt+1; p=lt+2; di=lt+3; remark=lt+4
    for item in range(1,101):
        col=item+2; letter=get_column_letter(col)
        if item<=nitems:
            ws.cell(ht,col).value=f'=COUNTIF({letter}{hs}:{letter}{ht-1},1)'
            ws.cell(ph,col).value=f'=IFERROR({letter}{ht}/COUNTA($B$'+str(hs)+':$B$'+str(ht-1)+'),0)'
            ws.cell(lt,col).value=f'=COUNTIF({letter}{ls}:{letter}{lt-1},1)'
            ws.cell(pl,col).value=f'=IFERROR({letter}{lt}/COUNTA($B$'+str(ls)+':$B$'+str(lt-1)+'),0)'
            ws.cell(p,col).value=f'=({letter}{ph}+{letter}{pl})/2'
            ws.cell(di,col).value=f'={letter}{ph}-{letter}{pl}'
            ws.cell(remark,col).value=f'=IF({letter}{di}>0.2,"G",IF({letter}{di}>0.1,"M","P"))'
        else:
            for row in (ht,ph,lt,pl,p,di,remark): ws.cell(row,col).value=None

def _workbook(template,out,analysis,m):
    wb=load_workbook(template); ws=wb["Items Analysis"]
    hn=max(16,len(analysis.high_group)); ln=max(16,len(analysis.low_group))
    et,el=hn-16,ln-16; _expand(ws,et,el)
    hs=18; ls=45+et
    _metadata(ws,m,len(analysis.data.records))
    _student_header(ws,16)
    _student_header(ws,ls-2)
    _group(ws,analysis.high_group,hs,hn,analysis.data.question_count,1)
    _group(ws,analysis.low_group,ls,ln,analysis.data.question_count,len(analysis.high_group)+1)
    _formulas(ws,analysis.data.question_count,hs,hn,ls,ln)
    wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True; wb.calculation.calcMode="auto"
    wb.save(out)

def _replace_span(paragraph,start,end,value):
    offsets=[]
    position=0
    for run in paragraph.runs:
        offsets.append((position,position+len(run.text)))
        position+=len(run.text)
    if not offsets:
        if start==0 and end==0:
            paragraph.add_run(value)
        return
    first=next((i for i,(_,stop) in enumerate(offsets) if stop>start),len(offsets)-1)
    last=next((i for i,(_,stop) in enumerate(offsets) if stop>=end),first)
    first_start=offsets[first][0]
    last_start=offsets[last][0]
    first_text=paragraph.runs[first].text[:start-first_start]
    last_text=paragraph.runs[last].text[end-last_start:]
    if first==last:
        paragraph.runs[first].text=first_text+value+last_text
        return
    paragraph.runs[first].text=first_text+value
    for index in range(first+1,last):
        paragraph.runs[index].text=""
    paragraph.runs[last].text=last_text


def _replace_pattern(paragraph,pattern,value,flags=0):
    for match in reversed(list(re.finditer(pattern,paragraph.text,flags))):
        _replace_span(paragraph,match.start(),match.end(),str(value))


def _report(path,analysis,m):
    if not SUMMARY_TEMPLATE.is_file():
        raise OutputError(f"Word summary template not found: {SUMMARY_TEMPLATE}")
    doc=Document(SUMMARY_TEMPLATE)
    total=len(analysis.items)
    counts={category:sum(item.classification==category for item in analysis.items)
            for category in ("Good","Marginal","Poor")}
    percentages={category:(counts[category]/total*100 if total else 0)
                 for category in counts}
    if counts["Good"]>total/2:
        good_opening="Most of the questions were found to be GOOD questions"
    elif counts["Good"]:
        good_opening="Some of the questions were found to be GOOD questions"
    else:
        good_opening="No questions were found to be GOOD questions"
    program=m.get("program","")
    program_name=PROGRAM_LONG_NAMES.get(program,program)
    exam_names={"Prelims":"preliminary examination","Midterms":"midterm examination",
                "Finals":"final examination"}
    number_pattern=r"(?:\{\{number_of_items\}\}|\{\{number_of_items[^}]*\})"
    for paragraph in doc.paragraphs:
        _replace_pattern(paragraph,r"Most of the questions were found to be GOOD questions",good_opening)
        _replace_pattern(paragraph,r"\{\{Program\}\}",program_name)
        _replace_pattern(paragraph,r"\{\{Subject\}\}",m.get("subject",""))
        _replace_pattern(paragraph,r"\{\{Prepared_By\}\}",m.get("prepared_by",""))
        _replace_pattern(paragraph,r"\{\{Department_Chairperson\}\}",m.get("department_chairperson",""))
        _replace_pattern(paragraph,r"GOOD/"+number_pattern,f"{counts['Good']}/{total}")
        _replace_pattern(paragraph,r"POOR/"+number_pattern,f"{counts['Poor']}/{total}")
        _replace_pattern(paragraph,r"\{\{%_poor\}\}",f"{percentages['Poor']:.1f}%")
        _replace_pattern(paragraph,r"%_GOOD",f"{percentages['Good']:.1f}%")
        _replace_pattern(paragraph,r"\{\{number_of_marginal\}\}",counts["Marginal"])
        _replace_pattern(paragraph,number_pattern,total)
        _replace_pattern(paragraph,r"final examination",exam_names.get(m.get("exam_type"),"examination"),re.I)
        _replace_pattern(paragraph,r"Department Chair, Biology Program",
                         f"Department Chair, {program} Program" if program else "Department Chair")
        _replace_pattern(paragraph,"\u2019", "'")

    details_title=doc.add_paragraph()
    details_title.paragraph_format.page_break_before=True
    details_title.paragraph_format.keep_with_next=True
    title_run=details_title.add_run("ITEM ANALYSIS DETAILS")
    title_run.bold=True
    title_run.font.name="Arial"
    title_run.font.size=Pt(12)
    for category in ("Good","Marginal","Poor"):
        items=[item for item in analysis.items if item.classification==category]
        heading=doc.add_paragraph()
        heading.paragraph_format.keep_with_next=True
        heading_run=heading.add_run(f"{category} Items ({len(items)})")
        heading_run.bold=True
        heading_run.font.name="Arial"
        heading_run.font.size=Pt(11)
        if not items:
            empty=doc.add_paragraph("None.")
            empty.paragraph_format.keep_with_next=True
            continue
        table=doc.add_table(rows=1,cols=3)
        table.autofit=False
        borders=OxmlElement("w:tblBorders")
        for edge in ("top","left","bottom","right","insideH","insideV"):
            border=OxmlElement(f"w:{edge}")
            border.set(qn("w:val"),"single")
            border.set(qn("w:sz"),"4")
            border.set(qn("w:space"),"0")
            border.set(qn("w:color"),"808080")
            borders.append(border)
        table._tbl.tblPr.append(borders)
        widths=(Inches(1.0),Inches(2.5),Inches(2.5))
        for column,width in zip(table.columns,widths):
            column.width=width
        headers=("Item","Difficulty (p)","Discrimination (D)")
        for cell,label in zip(table.rows[0].cells,headers):
            cell.width=widths[headers.index(label)]
            cell.text=label
            for run in cell.paragraphs[0].runs:
                run.bold=True
        repeat=OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"),"true")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for item in items:
            cells=table.add_row().cells
            for cell,width,value in zip(cells,widths,(str(item.number),f"{item.difficulty:.3f}",f"{item.discrimination:.3f}")):
                cell.width=width
                cell.text=value
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name="Arial"
                        run.font.size=Pt(10)
    for paragraph in doc.paragraphs[21:37]:
        paragraph.paragraph_format.keep_together=True
        paragraph.paragraph_format.keep_with_next=True
    doc.save(path)

def output_paths(output_dir,metadata,source_label=None):
    folder=Path(output_dir)
    stem=f"Item_Analysis_{safe_part(metadata['subject'],'Subject')}_{safe_part(metadata['exam_type'],'Exam')}"
    if source_label:
        stem+=f"_{safe_part(source_label,'Class')}"
    return folder/f"{stem}.xlsx",folder/f"{stem}_Summary.docx"

def generate_outputs(template_path,output_dir,analysis,metadata,source_label=None):
    template=Path(template_path)
    if not template.is_file(): raise OutputError(f"Excel template not found: {template}")
    folder=Path(output_dir); folder.mkdir(parents=True,exist_ok=True)
    if not SUMMARY_TEMPLATE.is_file(): raise OutputError(f"Word summary template not found: {SUMMARY_TEMPLATE}")
    xlsx,docx=output_paths(folder,metadata,source_label)
    _workbook(template,xlsx,analysis,metadata)
    report_metadata=dict(metadata)
    if source_label: report_metadata["class_label"]=source_label
    _report(docx,analysis,report_metadata)
    return xlsx,docx


