from pathlib import Path
import pytest
from docx import Document
from openpyxl import load_workbook
from item_analysis_maker.analysis import analyze, combine_datasets
from item_analysis_maker.parsing import InputData, InputFormatError, StudentRecord, parse_csv, parse_csv_batch
from item_analysis_maker.outputs import generate_outputs

ROOT = Path(__file__).resolve().parent

def _write_lms(path, questions, attempts):
    headers = ["Last name", "First name", "Email address", "Status", "Grade/100.00"]
    headers += [f"Q. {n} /1.00" for n in range(1, questions + 1)]
    lines = [",".join(headers)]
    for status, grade, answers in attempts:
        lines.append(",".join(["PrivateLast", "PrivateFirst", "private@example.test", status, str(grade)]
                             + [str(answer) for answer in answers]))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")

def _write_zipgrade(path, questions, attempts):
    headers = ["Quiz Name", "Class", "ZipGrade ID", "External Id", "First Name", "Last Name",
               "Num Questions", "Num Correct", "Percent Correct", "Key Version"]
    headers += [f"Q{n}" for n in range(1, questions + 1)]
    lines = [",".join(headers)]
    for class_name, score, answers in attempts:
        pct = score / questions * 100
        lines.append(",".join(["Quiz", class_name, "0", "", "PrivateFirst", "PrivateLast",
                               str(questions), str(score), str(pct), "A"] + [str(a) for a in answers]))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")

def test_lms_parser_filters_unfinished_and_drops_identity_fields(tmp_path):
    path = tmp_path / "class.csv"
    _write_lms(path, 2, [("Finished", 2, (1, 1)), ("In progress", 0, (0, 0))])
    data = parse_csv(path)
    assert data.format == "LMS" and data.question_count == 2
    assert len(data.records) == 1 and data.excluded_rows == 1
    assert not hasattr(data.records[0], "name")

def test_zipgrade_parser_uses_declared_item_count(tmp_path):
    path = tmp_path / "class.csv"
    _write_zipgrade(path, 3, [("A", 2, (1, 1, 0)), ("A", 1, (0, 1, 0))])
    data = parse_csv(path)
    assert data.format == "ZipGrade" and len(data.records) == 2
    assert data.question_count == 3 and len(data.records[0].answers) == 3

def test_ties_and_classification_thresholds():
    tied = (StudentRecord(100, (True,)), StudentRecord(90, (True,)), StudentRecord(90, (True,)),
            StudentRecord(80, (False,)), StudentRecord(70, (False,)))
    grouped = analyze(InputData("ZipGrade", tied, 1, 0))
    assert grouped.nominal_group_size == 2 and len(grouped.high_group) == 3 and len(grouped.low_group) == 2
    records = tuple(StudentRecord(score, (37-score < 3, 37-score < 2, 37-score in (0, 10)))
                    for score in range(37, 0, -1))
    result = analyze(InputData("ZipGrade", records, 3, 0))
    assert [item.classification for item in result.items] == ["Good", "Marginal", "Poor"]
    assert result.items[1].discrimination == pytest.approx(0.2)
    assert result.items[2].discrimination == pytest.approx(0.1)

def test_invalid_files_and_missing_headers(tmp_path):
    with pytest.raises(InputFormatError, match="\\.csv"):
        parse_csv(ROOT / "README.md")
    path = tmp_path / "bad.csv"
    path.write_text("name,score\nA,1\n", encoding="utf-8")
    with pytest.raises(InputFormatError, match="not recognized"):
        parse_csv(path)

def test_bad_score_and_no_finished(tmp_path):
    path = tmp_path / "bad.csv"
    path.write_text("Last name,First name,Email address,Status,Grade/100.00,Q. 1 /1.00\nPrivate,Person,p@example.test,Finished,nope,1\n", encoding="utf-8")
    with pytest.raises(InputFormatError, match="Invalid number"):
        parse_csv(path)
    path.write_text("Last name,First name,Email address,Status,Grade/100.00,Q. 1 /1.00\nPrivate,Person,p@example.test,In progress,1,1\n", encoding="utf-8")
    with pytest.raises(InputFormatError, match="No Finished"):
        parse_csv(path)

def test_outputs_are_anonymized_for_lms_and_zipgrade(tmp_path):
    lms = tmp_path / "lms.csv"
    zipgrade = tmp_path / "zip.csv"
    _write_lms(lms, 2, [("Finished", 2, (1, 1)), ("Finished", 1, (1, 0)), ("Finished", 0, (0, 0))])
    _write_zipgrade(zipgrade, 2, [("A", 2, (1, 1)), ("A", 1, (1, 0)), ("A", 0, (0, 0))])
    metadata = {"exam_type":"Prelims", "academic_year":"2026-2027", "semester":"1st Sem",
                "subject":"Sample Subject", "description":"Unit exam", "prepared_by":"Faculty A",
                "department_chairperson":"Chair B"}
    for path in (lms, zipgrade):
        data = parse_csv(path)
        xlsx, docx = generate_outputs(ROOT / "ITEM_ANALYSIS_FORMAT.xlsx", tmp_path / path.stem,
                                      analyze(data), metadata)
        sheet = load_workbook(xlsx, data_only=False)["Items Analysis"]
        assert sheet["E10"].value == "2026-2027" and sheet["J10"].value == "Sample Subject"
        assert sheet["X10"].value == len(data.records)
        assert sheet["A18"].value == 1 and sheet["B18"].value == "Student 1"
        assert sheet["B16"].value == "Student Names" and sheet["A10"].value.startswith("X ")
        all_cells = " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None)
        assert "PrivateLast" not in all_cells and "private@example.test" not in all_cells
        report = Document(docx)
        text = "\n".join(p.text for p in report.paragraphs)
        text += "\n" + "\n".join(c.text for table in report.tables for row in table.rows for c in row.cells)
        assert all(f"{category} Items" in text for category in ("Good", "Marginal", "Poor"))
        assert "PrivateLast" not in text and "private@example.test" not in text

def test_groups_are_capped_and_numbered_continuously(tmp_path):
    analysis = analyze(InputData("ZipGrade", tuple(StudentRecord(1, (True,)) for _ in range(50)), 1, 0))
    assert len(analysis.high_group) == 16 and len(analysis.low_group) == 16
    metadata = {"exam_type":"Finals", "academic_year":"2026-2027", "semester":"2nd Sem",
                "subject":"Tie", "description":"", "prepared_by":"", "department_chairperson":""}
    xlsx, _ = generate_outputs(ROOT / "ITEM_ANALYSIS_FORMAT.xlsx", tmp_path, analysis, metadata)
    sheet = load_workbook(xlsx, data_only=False)["Items Analysis"]
    assert sheet["A18"].value == 1 and sheet["B18"].value == "Student 1"
    assert sheet["A33"].value == 16 and sheet["B33"].value == "Student 16"
    assert sheet["A45"].value == 17 and sheet["B45"].value == "Student 17"
    assert sheet["A60"].value == 32 and sheet["B60"].value == "Student 32"
    assert sheet["X14"].value == "=MIN(ROUNDUP(X11,0),16)"
    assert sheet["B18"].font.name == "Arial Narrow" and sheet["B18"].font.sz == 10
    assert sheet["A45"].font.name == "Arial Narrow" and sheet["A45"].font.sz == 10
    assert sheet["B35"].value == "pH" and sheet["B62"].value == "pL"

def test_low_group_numbering_continues_after_actual_high_group_size(tmp_path):
    records = tuple(StudentRecord(score, (score >= 9,)) for score in range(16, 0, -1))
    analysis = analyze(InputData("ZipGrade", records, 1, 0))
    assert len(analysis.high_group) == 5 and len(analysis.low_group) == 5
    metadata = {"exam_type":"Midterms", "academic_year":"2026-2027", "semester":"1st Sem",
                "subject":"Numbering", "description":"", "prepared_by":"", "department_chairperson":""}
    xlsx, _ = generate_outputs(ROOT / "ITEM_ANALYSIS_FORMAT.xlsx", tmp_path, analysis, metadata)
    sheet = load_workbook(xlsx, data_only=False)["Items Analysis"]
    assert [sheet.cell(row, 2).value for row in range(18, 23)] == [f"Student {n}" for n in range(1, 6)]
    assert [sheet.cell(row, 2).value for row in range(45, 50)] == [f"Student {n}" for n in range(6, 11)]
    assert sheet["A45"].value == 6 and sheet["A49"].value == 10

def test_csv_batch_rejects_different_item_counts(tmp_path):
    two_items = tmp_path / "two.csv"
    one_item = tmp_path / "one.csv"
    _write_lms(two_items, 2, [("Finished", 2, (1, 1))])
    _write_lms(one_item, 1, [("Finished", 1, (1,))])
    with pytest.raises(InputFormatError, match="same number of items"):
        parse_csv_batch((two_items, one_item))

def test_same_course_csvs_pool_for_one_group_pair(tmp_path):
    first = InputData("ZipGrade", tuple(StudentRecord(score, (score >= 9,)) for score in range(8, 0, -1)), 1, 0)
    second = InputData("ZipGrade", tuple(StudentRecord(score + 8, (score >= 5,)) for score in range(8, 0, -1)), 1, 0)
    combined = combine_datasets((first, second))
    result = analyze(combined)
    assert len(combined.records) == 16
    assert len(result.high_group) == 5 and len(result.low_group) == 5
    metadata = {"exam_type":"Prelims", "academic_year":"2026-2027", "semester":"1st Sem",
                "subject":"Combined Course", "description":"", "prepared_by":"", "department_chairperson":""}
    xlsx, docx = generate_outputs(ROOT / "ITEM_ANALYSIS_FORMAT.xlsx", tmp_path, result, metadata)
    assert xlsx.exists() and docx.exists()
    assert len(list(tmp_path.glob("*.xlsx"))) == 1
    assert len(list(tmp_path.glob("*.docx"))) == 1
