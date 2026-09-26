from io import BytesIO
import zipfile
from docx import Document
from fastapi.testclient import TestClient
from item_analysis_maker.web.app import app

client=TestClient(app)

def _metadata():
    return {"exam_type":"Prelims","academic_year":"2026-2027","semester":"1st Sem",
            "subject":"Web Course","description":"Combined class exports",
            "prepared_by":"Faculty","department_chairperson":"Chair"}

def _zip_csv(item_count, attempts):
    headers=["Quiz Name","Class","ZipGrade ID","External Id","First Name","Last Name",
             "Num Questions","Num Correct","Percent Correct","Key Version"]
    headers += [f"Q{i}" for i in range(1,item_count+1)]
    lines=[",".join(headers)]
    for group,score,answers in attempts:
        pct=score/item_count*100
        lines.append(",".join(["Quiz",group,"0","","Synthetic","Learner",str(item_count),
                               str(score),str(pct),"A"]+[str(value) for value in answers]))
    return "\n".join(lines)+"\n"

def test_homepage_renders_multi_csv_form():
    response=client.get("/")
    assert response.status_code==200
    assert 'href="/static/style.css"' in response.text
    assert 'name="files"' in response.text
    assert "multiple" in response.text
    assert "same item count" in response.text
    assert 'name="program"' in response.text
    assert 'id="program-long-name"' in response.text and "readonly" in response.text
    assert "Pharmacy — Mr. Aaron Dell A. Cobeng" in response.text

def test_static_stylesheet_is_served_from_root_relative_path():
    response=client.get("/static/style.css")
    assert response.status_code==200
    assert ".card" in response.text

def test_validate_and_generate_reject_mismatched_item_counts():
    two=_zip_csv(2,[("A",2,(1,1))]).encode()
    one=_zip_csv(1,[("B",1,(1,))]).encode()
    fields=[("files",("Class_Two.csv",two,"text/csv")),("files",("Class_One.csv",one,"text/csv"))]
    response=client.post("/validate",data=_metadata(),files=fields)
    assert response.status_code==200
    assert "same number of items" in response.text and "No files were generated" in response.text
    fields=[("files",("Class_Two.csv",two,"text/csv")),("files",("Class_One.csv",one,"text/csv"))]
    response=client.post("/generate",data=_metadata(),files=fields)
    assert response.status_code==200
    assert response.headers["content-type"].startswith("text/html")
    assert "same number of items" in response.text

def test_generate_rejects_unknown_program():
    fields=[("files",("Class_A.csv",_zip_csv(1,[("A",1,(1,))]).encode(),"text/csv"))]
    metadata=_metadata()
    metadata["program"]="Unknown Program"
    response=client.post("/generate",data=metadata,files=fields)
    assert response.status_code==200
    assert "valid program" in response.text

def test_generate_returns_one_archive_for_pooled_csvs():
    first=_zip_csv(1,[("A",1,(1,)),("A",0,(0,))]).encode()
    second=_zip_csv(1,[("B",1,(1,)),("B",0,(0,))]).encode()
    files=[("files",("Class_A.csv",first,"text/csv")),("files",("Class_B.csv",second,"text/csv"))]
    metadata=_metadata()
    metadata.update(program="Pharmacy",department_chairperson="")
    response=client.post("/generate",data=metadata,files=files)
    assert response.status_code==200
    assert response.headers["content-type"].startswith("application/zip")
    with zipfile.ZipFile(BytesIO(response.content)) as archive:
        names=archive.namelist()
        assert len(names)==2
        workbook_name=next(name for name in names if name.endswith(".xlsx"))
        assert "xl/media/image1.png" in zipfile.ZipFile(BytesIO(archive.read(workbook_name))).namelist()
        report_name=next(name for name in names if name.endswith(".docx"))
        report=Document(BytesIO(archive.read(report_name)))
        report_text="\n".join(paragraph.text for paragraph in report.paragraphs)
        report_text+="\n"+"\n".join(cell.text for table in report.tables for row in table.rows for cell in row.cells)
        assert "Diploma in Pharmacy Assisting Leading to" in report_text
        assert "Mr. Aaron Dell A. Cobeng, RPh, MA ELM" in report_text
        assert "Academic Year:" not in report_text and "Semester:" not in report_text
        assert "Students analyzed:" not in report_text and "High group:" not in report_text
        assert "Good Items (" in report_text and "Difficulty (p)" in report_text and "Discrimination (D)" in report_text
        assert len(report.tables) == 1
        assert b"PrivateFirst" not in b"".join(archive.read(name) for name in names)
